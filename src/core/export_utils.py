"""
Export Utilities for PDF and DOCX
"""

import os
from pathlib import Path
from typing import Optional
from datetime import datetime

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
    from reportlab.lib import colors
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class ExportUtils:
    """Utilities for exporting content to PDF and DOCX formats"""
    
    @staticmethod
    def export_to_pdf(content: dict, output_path: str) -> bool:
        """
        Export content to PDF
        
        Args:
            content: Dictionary with 'title', 'body', and optional 'data'
            output_path: Path to save the PDF file
        
        Returns:
            True if successful, False otherwise
        """
        if not REPORTLAB_AVAILABLE:
            print("ReportLab not installed. PDF export unavailable.")
            return False
        
        try:
            doc = SimpleDocTemplate(output_path, pagesize=letter)
            story = []
            styles = getSampleStyleSheet()
            
            # Title
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#007AFF'),
                spaceAfter=30,
                alignment=1  # Center
            )
            
            title = Paragraph(content.get('title', 'RecursiveLearn Export'), title_style)
            story.append(title)
            story.append(Spacer(1, 0.2 * inch))
            
            # Metadata
            date_style = styles['Normal']
            date_text = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            story.append(Paragraph(date_text, date_style))
            story.append(Spacer(1, 0.3 * inch))
            
            # Body content
            body_style = styles['BodyText']
            body = content.get('body', '')
            
            # Split by lines and add paragraphs
            for line in body.split('\n'):
                if line.strip():
                    p = Paragraph(line, body_style)
                    story.append(p)
                    story.append(Spacer(1, 0.1 * inch))
            
            # Data table (if provided)
            if 'data' in content and content['data']:
                story.append(Spacer(1, 0.3 * inch))
                
                data = content['data']
                t = Table(data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#007AFF')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                story.append(t)
            
            # Build PDF
            doc.build(story)
            return True
            
        except Exception as e:
            print(f"Error exporting to PDF: {e}")
            return False
    
    @staticmethod
    def export_to_docx(content: dict, output_path: str) -> bool:
        """
        Export content to DOCX
        
        Args:
            content: Dictionary with 'title', 'body', and optional 'data'
            output_path: Path to save the DOCX file
        
        Returns:
            True if successful, False otherwise
        """
        if not DOCX_AVAILABLE:
            print("python-docx not installed. DOCX export unavailable.")
            return False
        
        try:
            doc = Document()
            
            # Title
            title = doc.add_heading(content.get('title', 'RecursiveLearn Export'), 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Metadata
            date_para = doc.add_paragraph()
            date_run = date_para.add_run(
                f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            )
            date_run.italic = True
            date_run.font.size = Pt(10)
            date_run.font.color.rgb = RGBColor(128, 128, 128)
            
            doc.add_paragraph()  # Spacer
            
            # Body content
            body = content.get('body', '')
            for line in body.split('\n'):
                if line.strip():
                    doc.add_paragraph(line)
            
            # Data table (if provided)
            if 'data' in content and content['data']:
                doc.add_paragraph()  # Spacer
                
                data = content['data']
                if data:
                    table = doc.add_table(rows=len(data), cols=len(data[0]))
                    table.style = 'Light Grid Accent 1'
                    
                    for i, row in enumerate(data):
                        for j, cell in enumerate(row):
                            table.cell(i, j).text = str(cell)
            
            # Save
            doc.save(output_path)
            return True
            
        except Exception as e:
            print(f"Error exporting to DOCX: {e}")
            return False
    
    @staticmethod
    def export_problem_solution(recurrence: str, initial_conditions: dict,
                               solution: dict, output_path: str,
                               format: str = 'pdf') -> bool:
        """
        Export a problem and its solution
        
        Args:
            recurrence: Recurrence relation string
            initial_conditions: Initial conditions dict
            solution: Solution dictionary from solver
            output_path: Output file path
            format: 'pdf' or 'docx'
        
        Returns:
            True if successful
        """
        content = {
            'title': 'Recursive Relation Solution',
            'body': f"Recurrence Relation:\n{recurrence}\n\n"
                   f"Initial Conditions:\n{', '.join([f'a({n}) = {v}' for n, v in sorted(initial_conditions.items())])}\n\n"
                   f"Solution Steps:\n\n"
        }
        
        # Add solution steps
        if 'steps' in solution:
            for i, step in enumerate(solution['steps'], 1):
                content['body'] += f"Step {i}: {step['title']}\n{step['content']}\n\n"
        
        # Add final solution
        if 'solution' in solution:
            content['body'] += f"\nFinal Solution:\n{solution['solution']}"
        
        if format.lower() == 'pdf':
            return ExportUtils.export_to_pdf(content, output_path)
        elif format.lower() == 'docx':
            return ExportUtils.export_to_docx(content, output_path)
        else:
            return False
